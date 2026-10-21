"""
Tests del filtro de estado_linea (defensa frente a grupos disfrazados de estado).

La guia ERP (SKU_BX.xlsx) define 4 estados de linea: NUEVO, IMPORTADO,
TRADICIONAL y NACIONAL. El API de stock, cuando el catalogo no trae estado,
lo rellena con el GRUPO del producto: hay que blindar el dato.
"""

import pandas as pd

from src.core.ingest import Ingester


def _df_estado(estado="TAJADORES", sku="79026"):
    return pd.DataFrame(
        {
            "ID_PEDIDO": ["P1"],
            "SKU": [sku],
            "NOM_ARTICULO": ["TIJERA VINIFAN PANDA"],
            "linea": ["79"],
            "grupo": ["02"],
            "categoria": ["VINIFAN"],
            "estado_linea": [estado],
            "CAN_PEDIDO": [100],
            "CAN_DESPACHO": [0],
            "PEDIDO_IMPORTE_ARITCULO": [177.0],
            "PRE_ARTICULO": [177.0],
            "VIGENCIA": [100.0],
            "CONDICION_PAGO": ["ANTICIPO EN LETRAS"],
            "PEDIDO_ESTADO": ["APROBADO"],
        }
    )


def test_estado_linea_ajeno_se_limpia():
    """Un grupo (TAJADORES) no puede quedar como estado de linea."""
    ing = Ingester()
    out = ing.compute_derived(_df_estado("TAJADORES"))
    assert str(out.loc[0, "estado_linea"]).strip() == ""
    # el valor original se conserva para auditoria
    assert str(out.loc[0, "estado_linea_raw"]).strip() == "TAJADORES"
    assert ing._estados_linea_invalidos == 1


def test_estado_linea_valido_se_conserva():
    for estado in ("NUEVO", "IMPORTADO", "TRADICIONAL", "NACIONAL"):
        ing = Ingester()
        out = ing.compute_derived(_df_estado(estado))
        assert str(out.loc[0, "estado_linea"]).strip() == estado
        assert ing._estados_linea_invalidos == 0


def test_reporte_reporta_los_estados_fuera_de_guia(tmp_path):
    """El descarte se informa: no se pierde en silencio."""
    from openpyxl import load_workbook

    from src.core.exporters import export_reporte_acciones

    ing = Ingester()
    df = ing.compute_derived(_df_estado("RESALTADORES"))
    path = export_reporte_acciones(
        df,
        familia="ANTICIPO",
        output_path=str(tmp_path / "r.xlsx"),
        guardar_foto=False,
    )
    wb = load_workbook(path)
    panel = " ".join(str(c.value) for r in wb["PANEL"].iter_rows() for c in r if c.value)
    assert "Estados de linea fuera de la guia" in panel
    assert "RESALTADORES" in panel
