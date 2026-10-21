"""Tests Fase 1 + Fase 2: cache sin red, EXCEPCIONES y alertas 121."""

import sys
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.settings import UMBRAL_CAJAS_121_OBSERVACION  # noqa: E402


# ---------------------------------------------------------------- helpers
def _df_stock(**kwargs):
    """DataFrame minimo con las columnas que pide el pipeline de stock."""
    base = {
        "ID_PEDIDO": ["P1", "P1"],
        "ID_ARTICULO": ["AAA111", "BBB222"],
        "SKU": ["AAA111", "BBB222"],
        "NOM_ARTICULO": ["PROD A", "PROD B"],
        "ID_VENDEDOR": ["01", "02"],
        "NOM_VENDEDOR": ["V1", "V2"],
        "COD_CLIENTE": ["C1", "C2"],
        "NOM_CLIENTE": ["CLI1", "CLI2"],
        "CONDICION_PAGO": ["ANTICIPO EN LETRAS", "ANTICIPO EN LETRAS"],
        "PEDIDO_ESTADO": ["APROBADO", "APROBADO"],
        "CAN_PEDIDO": [100, 50],
        "CAN_DESPACHO": [0, 0],
        "PRE_ARTICULO": [10.0, 20.0],
        "PEDIDO_IMPORTE_ARITCULO": [1000.0, 1000.0],
        "VIGENCIA": [100, 50],
        "linea": ["78 - ARCHIVO", "79 - ACCESORIOS"],
        "grupo": ["FOLDER", "BORRADOR"],
        "un_bx": [10, 12],
        "stock_api_disponible": [{"VES": 0, "121": 5000}, {"VES": 0, "121": 6}],
    }
    base.update(kwargs)
    return pd.DataFrame(base)


# ====================================================================
# FASE 1: el cache se aplica aunque este vencido (--sin-stock)
# ====================================================================
class TestCacheSinRed:
    def test_apply_stock_cache_clasifica_121(self, tmp_path, monkeypatch):
        """Sin red, el cache local igual clasifica VES/121/otros."""
        import src.core.ingest as ing_mod

        ing = ing_mod.Ingester()
        df = _df_stock()

        # Cache con 121 = 5000 para AAA111
        ing._stock_cache = {"AAA111": {"VES": 0, "121": 5000}, "BBB222": {"VES": 0, "121": 6}}
        ing._render_catalog = {}

        out = ing._apply_stock_cache(df)

        assert "stock_api_121" in out.columns
        assert float(out.loc[out["SKU"] == "AAA111", "stock_api_121"].iloc[0]) == 5000
        assert float(out.loc[out["SKU"] == "BBB222", "stock_api_121"].iloc[0]) == 6
        # VES = 0 en ambos
        assert float(out.loc[out["SKU"] == "AAA111", "stock_api_ves"].iloc[0]) == 0

    def test_load_stock_cache_ignore_ttl(self, tmp_path, monkeypatch):
        """Con ignore_ttl=True se lee el cache aunque este vencido."""
        import json
        import time

        import src.config.settings as st
        import src.core.ingest as ing_mod

        cache_dir = tmp_path / ".cache"
        cache_dir.mkdir()
        (cache_dir / "stock_cache.json").write_text(
            json.dumps(
                {
                    "ts": time.time() - 99999,  # muy viejo
                    "stock": {"AAA111": {"VES": 3, "121": 7}},
                    "catalog": {},
                }
            ),
            encoding="utf-8",
        )
        monkeypatch.setattr(ing_mod, "DATA_DIR", cache_dir)

        ing = ing_mod.Ingester()
        ing._stock_cache = None
        ing._render_catalog = {}

        # Vencido (TTL 15 min) -> sin ignore_ttl no hay nada
        assert ing._load_stock_cache() == {}

        # Con ignore_ttl si hay datos
        ing._stock_cache = None
        ing._render_catalog = {}
        cache = ing._load_stock_cache(ignore_ttl=True)
        assert cache.get("AAA111", {}).get("121") == 7

    def test_apply_stock_cache_sin_cache_no_rompe(self):
        """Sin cache disponible el pipeline sigue (stock 0, sin excepcion)."""
        import src.core.ingest as ing_mod

        ing = ing_mod.Ingester()
        ing._stock_cache = {}
        ing._render_catalog = {}

        out = ing._apply_stock_cache(_df_stock())
        assert len(out) == 2
        assert float(out["stock_api_121"].sum()) == 0


# ====================================================================
# FASE 1: hoja EXCEPCIONES
# ====================================================================
class TestHojaExcepciones:
    def test_excepciones_agrupa_por_tipo(self, tmp_path):
        """La hoja EXCEPCIONES separa REVISION_MANUAL / FUERA_DE_ALLOWLIST."""
        from openpyxl import Workbook

        from src.core.exporters import _write_excepciones

        wb = Workbook()
        wb.remove(wb.active)

        df = pd.DataFrame(
            [
                {
                    "ID_ARTICULO": "SINLINEA1",
                    "SKU": "SINLINEA1",
                    "NOM_ARTICULO": "PRODUCTO SIN LINEA",
                    "linea": "",
                    "categoria": "VINIFAN",
                    "origen_linea": "VACIA",
                    "saldo": 100.0,
                    "valor_saldo": 500.0,
                    "ID_PEDIDO": "P9",
                    "ID_VENDEDOR": "05",
                    "NOM_VENDEDOR": "V5",
                    "COD_CLIENTE": "C9",
                    "NOM_CLIENTE": "CLI9",
                    "es_linea_valida": False,
                }
            ]
        )
        rechazados = pd.DataFrame(
            [
                {
                    "ID_ARTICULO": "FUERA1",
                    "SKU": "FUERA1",
                    "NOM_ARTICULO": "PRODUCTO FUERA",
                    "linea": "20",
                    "categoria": "INDUSTRIAL",
                    "origen_linea": "CODIGO",
                    "saldo": 50.0,
                    "valor_saldo": 250.0,
                    "ID_PEDIDO": "P8",
                    "motivo_rechazo": "LINEA_FUERA_DE_ALLOWLIST",
                }
            ]
        )

        _write_excepciones(wb, df, rechazados, [])
        ws = wb["EXCEPCIONES"]
        tipos = [ws.cell(row=r, column=1).value for r in range(4, ws.max_row + 1)]
        assert "REVISION_MANUAL" in tipos
        assert "FUERA_DE_ALLOWLIST" in tipos

    def test_skus_colados_detecta_interseccion(self):
        """Un SKU en aceptado Y rechazado es un colado (fallo de integridad)."""
        from src.core.exporters import _skus_colados

        df = pd.DataFrame({"SKU": ["A1", "A2"]})
        rej = pd.DataFrame({"SKU": ["A2", "A3"]})
        colados = _skus_colados(df, rej)
        assert colados == ["A2"]

    def test_skus_colados_sin_interseccion(self):
        """Sin interseccion no hay colados (caso normal)."""
        from src.core.exporters import _skus_colados

        df = pd.DataFrame({"SKU": ["A1", "A2"]})
        rej = pd.DataFrame({"SKU": ["B1", "B2"]})
        assert _skus_colados(df, rej) == []


# ====================================================================
# FASE 2: alertas del almacen 121
# ====================================================================
class TestAlertas121:
    @staticmethod
    def _traslado(sku, saldo, s121, un_bx, ves=0):
        return {
            "SKU": sku,
            "ID_ARTICULO": sku,
            "NOM_ARTICULO": f"PROD {sku}",
            "ID_PEDIDO": "P1",
            "ID_VENDEDOR": "01",
            "NOM_VENDEDOR": "V",
            "COD_CLIENTE": "C",
            "NOM_CLIENTE": "CL",
            "CONDICION_PAGO": "ANTICIPO EN LETRAS",
            "PEDIDO_ESTADO": "APROBADO",
            "PEDIDO_FECHA": "2026-10-01",
            "linea": "78",
            "estado_linea": "NUEVO",
            "saldo": float(saldo),
            "un_bx": float(un_bx),
            "precio_unitario": 10.0,
            "stock_real": float(ves),
            "stock_api_ves": float(ves),
            "stock_api_121": float(s121),
        }

    def _exportar_traslado(self, tmp_path, rows):
        from openpyxl import Workbook

        from src.core.exporters import _write_traslado

        wb = Workbook()
        wb.remove(wb.active)
        _write_traslado(wb, pd.DataFrame(rows))
        return wb["TRASLADO_121"]

    def _leer_filas(self, ws):
        """Devuelve dicts de las filas de datos (desde la fila 5)."""
        headers = [ws.cell(row=3, column=c).value for c in range(1, ws.max_column + 1)]
        filas = []
        for r in range(5, ws.max_row + 1):
            vals = {
                headers[c - 1]: ws.cell(row=r, column=c).value for c in range(1, len(headers) + 1)
            }
            if vals.get("SKU"):
                filas.append(vals)
        return filas

    def test_121_una_caja_trasladable(self, tmp_path):
        """<= 1 caja en 121 -> TRASLADABLE (sin observacion)."""
        # 12 unidades con un_bx=12 -> 1 caja exacta
        ws = self._exportar_traslado(tmp_path, [self._traslado("SKU1", 100, 12, 12)])
        filas = self._leer_filas(ws)
        assert len(filas) == 1
        assert filas[0]["alerta_121"] == "TRASLADABLE"
        assert float(filas[0]["cajas_121"]) == 1.0

    def test_121_mas_de_una_caja_observacion(self, tmp_path):
        """> 1 caja en 121 -> TRASLADABLE + OBSERVACION (revisar conteo)."""
        ws = self._exportar_traslado(tmp_path, [self._traslado("SKU2", 100, 100, 10)])
        filas = self._leer_filas(ws)
        assert filas[0]["alerta_121"] == "TRASLADABLE + OBSERVACION"
        assert float(filas[0]["cajas_121"]) == 10.0

    def test_121_sobra_stock_observacion(self, tmp_path):
        """121 con stock pero SIN faltante -> OBSERVACION (sobra en QC)."""
        ws = self._exportar_traslado(tmp_path, [self._traslado("SKU3", 0, 500, 10)])
        filas = self._leer_filas(ws)
        assert filas[0]["alerta_121"] == "OBSERVACION"

    def test_121_usa_un_bx_para_las_cajas(self, tmp_path):
        """Las cajas se calculan con el factor de empaque (un_bx)."""
        # 100 und, un_bx=25 -> 4 cajas (mismo stock, distinto empaque)
        ws = self._exportar_traslado(tmp_path, [self._traslado("SKU4", 500, 100, 25)])
        filas = self._leer_filas(ws)
        assert float(filas[0]["cajas_121"]) == 4.0

    def test_121_columnas_de_numeros_claros(self, tmp_path):
        """La hoja expone debo / VES / 121 / cajas / a_trasladar."""
        ws = self._exportar_traslado(tmp_path, [self._traslado("SKU5", 70, 100, 10)])
        filas = self._leer_filas(ws)
        f = filas[0]
        assert float(f["saldo_total"]) == 70
        assert float(f["faltante_total"]) == 70
        assert float(f["stock_ves"]) == 0
        assert float(f["stock_121"]) == 100
        assert float(f["cajas_121"]) == 10.0
        assert float(f["a_trasladar"]) == 70  # min(debo, 121)

    def test_121_umbral_configurable(self):
        """El umbral de observacion se puede mover por entorno."""
        assert UMBRAL_CAJAS_121_OBSERVACION == 1.0
