"""
Test del par canonico de estados: estado_linea (producto) + estado_pedido (pedido).

- `estado_linea` es el nombre canonico (lo usan SKU_BX, master-data, el API de
  stock y nuestro pipeline). Sus valores validos son los 4 de la guia ERP.
- `estado_pedido` es el estado del PEDIDO (APROBADO / INGRESADO). En el df
  interno se conserva `PEDIDO_ESTADO` (nombre crudo del ERP y nombre de la
  columna en la DB); el libro lo expone como `estado_pedido`.
"""

import pandas as pd

from src.core.ingest import Ingester


def _df():
    return pd.DataFrame(
        {
            "ID_PEDIDO": ["P1", "P1"],
            "SKU": ["S1", "S2"],
            "NOM_ARTICULO": ["A", "B"],
            "linea": ["79", "78"],
            "grupo": ["02", "FOLDER"],
            "categoria": ["VINIFAN", "VINIFAN"],
            "estado_linea": ["NUEVO", "TRADICIONAL"],
            "CAN_PEDIDO": [10.0, 20.0],
            "CAN_DESPACHO": [0.0, 0.0],
            "PEDIDO_IMPORTE_ARITCULO": [100.0, 200.0],
            "PRE_ARTICULO": [10.0, 10.0],
            "VIGENCIA": [10.0, 20.0],
            "CONDICION_PAGO": ["ANTICIPO EN LETRAS", "ANTICIPO EN LETRAS"],
            "PEDIDO_ESTADO": ["APROBADO", "APROBADO"],
            "stock_api_disponible": [{"CENTRO LOGISTICO VES": 100}] * 2,
        }
    )


def test_el_df_conserva_el_nombre_crudo_del_erp():
    """PEDIDO_ESTADO no se renombra en el df: es tambien columna de la DB."""
    out = Ingester().compute_derived(_df())
    assert "PEDIDO_ESTADO" in out.columns
    assert str(out.loc[0, "PEDIDO_ESTADO"]) == "APROBADO"


def test_el_libro_expone_el_par_canonico(tmp_path):
    from openpyxl import load_workbook

    from src.core.exporters import export_reporte_acciones

    ing = Ingester()
    df = ing.compute_derived(_df())
    path = export_reporte_acciones(
        df,
        familia="ANTICIPO",
        output_path=str(tmp_path / "r.xlsx"),
        guardar_foto=False,
    )
    wb = load_workbook(path)
    cabeceras = [c.value for c in wb["PEDIDOS"][3]]
    assert "estado_pedido" in cabeceras
    assert "estado_linea" in cabeceras
    assert "PEDIDO_ESTADO" not in cabeceras
    assert "estados_linea" not in cabeceras
    # y los valores llegan bien
    filas = list(wb["PEDIDOS"].iter_rows(min_row=4, values_only=True))
    cab = [c.value for c in wb["PEDIDOS"][3]]
    i_est = cab.index("estado_linea")
    i_ped = cab.index("estado_pedido")
    assert filas[0][i_est] == "NUEVO, TRADICIONAL"
    assert filas[0][i_ped] == "APROBADO"


def test_pedidos_bloqueados_tienen_estado_linea(tmp_path):
    """El pedido bloqueado tambien informa sus estados de linea."""
    from src.core.analytics import PedidosAnalytics

    df = _df()
    df.loc[1, "PEDIDO_ESTADO"] = "INGRESADO"
    out = Ingester().compute_derived(df)
    bloq = PedidosAnalytics(out).get_pedidos_bloqueados()
    assert not bloq.empty
    assert "estado_linea" in bloq.columns
    assert str(bloq.iloc[0]["estado_linea"]) == "TRADICIONAL"
