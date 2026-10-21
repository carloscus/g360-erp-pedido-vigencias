"""
Tests del precio definitivo (sale del reporte de vigencias) y de su control.

Regla del ERP: PEDIDO_IMPORTE_ARITCULO = (PRE_ARTICULO / VIGENCIA) * CAN_PEDIDO
El precio unitario definitivo es PRE_ARTICULO / VIGENCIA, que es
algebraicamente igual a PEDIDO_IMPORTE_ARITCULO / CAN_PEDIDO.

El `precio` del catalogo (API / master-data) es REFERENCIAL (precio de lista)
y no se cruza con este dato.
"""

import pandas as pd
import pytest

from src.core.ingest import Ingester


def _df_precio(pre_art=5542.60, vigencia=39.59, can=60.0, importe=8400.00, can_desp=0.0):
    return pd.DataFrame(
        {
            "ID_PEDIDO": ["P1"],
            "SKU": ["S1"],
            "NOM_ARTICULO": ["ARTICULO"],
            "linea": ["78"],
            "grupo": ["01"],
            "categoria": ["VINIFAN"],
            "CAN_PEDIDO": [can],
            "CAN_DESPACHO": [can_desp],
            "PEDIDO_IMPORTE_ARITCULO": [importe],
            "PRE_ARTICULO": [pre_art],
            "VIGENCIA": [vigencia],
            "CONDICION_PAGO": ["ANTICIPO EN LETRAS"],
            "PEDIDO_ESTADO": ["APROBADO"],
        }
    )


def test_precio_unitario_es_el_de_la_vigencia():
    """precio_unitario == PRE_ARTICULO / VIGENCIA."""
    out = Ingester().compute_derived(_df_precio())
    # 5542.60 / 39.59 = 140.00  (y 8400 / 60 = 140.00: es la misma cifra)
    assert float(out.loc[0, "precio_unitario"]) == 140.0


def test_valor_saldo_usa_el_importe_del_erp():
    out = Ingester().compute_derived(_df_precio(can_desp=20.0))
    # saldo 40 de 60 pedidos -> 40/60 * 8400 = 5600
    assert float(out.loc[0, "valor_saldo"]) == 5600.0


def test_control_detecta_un_importe_incoherente():
    """Si el importe no cumple la formula, el control lo marca."""
    out = Ingester().compute_derived(_df_precio(importe=9999.0))
    assert bool(out.loc[0, "precio_verificable"]) is True
    assert bool(out.loc[0, "precio_coherente"]) is False
    assert float(out.loc[0, "precio_desvio"]) > 0.05


def test_control_acepta_un_importe_coherente():
    out = Ingester().compute_derived(_df_precio())
    assert bool(out.loc[0, "precio_verificable"]) is True
    assert bool(out.loc[0, "precio_coherente"]) is True
    assert float(out.loc[0, "precio_desvio"]) == 0.0


def test_lineas_sin_vigencia_no_son_verificables():
    """VIGENCIA = 0 -> PRE_ARTICULO = 0: no hay precio que contrastar."""
    df = _df_precio(pre_art=0.0, vigencia=0.0, importe=1500.0)
    out = Ingester().compute_derived(df)
    assert bool(out.loc[0, "precio_verificable"]) is False
    # no es un error: se marca coherente (no aplica) y no ensucia el KPI
    assert bool(out.loc[0, "precio_coherente"]) is True


def test_tolerancia_relativa_para_vigencia_minima():
    """Con VIGENCIA=0.01 la division amplifica el redondeo: debe tolerarse.

    PRE_ARTICULO=0.64, VIGENCIA=0.01, CAN=100.19 -> la formula da 6412.16
    y el ERP trae 6362.07: 0.79% de diferencia, dentro de la tolerancia.
    """
    df = _df_precio(pre_art=0.64, vigencia=0.01, can=100.19, importe=6362.07)
    out = Ingester().compute_derived(df)
    assert bool(out.loc[0, "precio_coherente"]) is True


def test_precio_unitario_por_pedido_es_ponderado():
    """El precio del pedido se pondera por unidades pendientes."""
    df = pd.DataFrame(
        {
            "ID_PEDIDO": ["P1", "P1"],
            "SKU": ["S1", "S2"],
            "NOM_ARTICULO": ["A", "B"],
            "linea": ["78", "78"],
            "grupo": ["01", "01"],
            "categoria": ["VINIFAN", "VINIFAN"],
            "CAN_PEDIDO": [10.0, 10.0],
            "CAN_DESPACHO": [0.0, 0.0],
            "PEDIDO_IMPORTE_ARITCULO": [100.0, 300.0],
            "PRE_ARTICULO": [10.0, 30.0],
            "VIGENCIA": [1.0, 1.0],
            "CONDICION_PAGO": ["ANTICIPO EN LETRAS"] * 2,
            "PEDIDO_ESTADO": ["APROBADO"] * 2,
        }
    )
    out = Ingester().compute_derived(df)
    from src.core.analytics import PedidosAnalytics

    acc = PedidosAnalytics(out).resumen_accion_pedido("ANTICIPO", 1200.0)
    # (10*10 + 30*10) / 20 = 20
    assert float(acc.iloc[0]["precio_unitario"]) == 20.0


def test_reporte_expone_precio_y_control(tmp_path):
    """El libro expone el precio unitario y el control de coherencia."""
    from openpyxl import load_workbook

    from src.core.exporters import export_reporte_acciones

    ing = Ingester()
    df = ing.compute_derived(_df_precio())
    path = export_reporte_acciones(
        df,
        familia="ANTICIPO",
        output_path=str(tmp_path / "r.xlsx"),
        guardar_foto=False,
    )
    wb = load_workbook(path)
    ped = wb["PEDIDOS"]
    cabeceras = [c.value for c in ped[3]]
    assert "precio_unitario" in cabeceras
    panel = " ".join(str(c.value) for r in wb["PANEL"].iter_rows() for c in r if c.value)
    assert "Coherencia de precios" in panel
